"""Apply the 2026-10-22 verification pass to the legacy workbook
financial/model/freight-recovery-financial-model.xlsx (superseded; kept for history).

Changes: turns three text cells that start with "=" (Excel reads them as broken
formulas) into plain text, marks the workbook superseded, and adds a Sources sheet.
Model values are not changed: this workbook is not the planning model.
"""
from pathlib import Path

import openpyxl
from openpyxl.styles import Font

ROOT = Path(__file__).resolve().parents[2]
WB = ROOT / "financial" / "model" / "freight-recovery-financial-model.xlsx"
ACCESSED = "2026-10-22"

wb = openpyxl.load_workbook(WB)
log = []

ue = wb["unit_economics"]
for ref in ("C3", "C4", "C10"):
    old = ue[ref].value
    if isinstance(old, str) and old.startswith("="):
        ue[ref].value = "calc: " + old.lstrip("= ").strip()
        log.append((f"unit_economics!{ref}", old, ue[ref].value, "text, not a formula"))

rv = wb["Source review 2026-10-19"]
rv["A15"] = "2026-10-22 update"
rv["B15"] = ("SUPERSEDED by business/investor-pack/scoup-financial-model.xlsx, the planning model "
             "quoted in the investor documents. Values here are kept for history and were not "
             "recalculated. Malformed '=' text cells in unit_economics fixed; Sources sheet added. "
             "Links to _launch/ point to an untracked file; use business/data-and-sources.md.")

if "Sources" in wb.sheetnames:
    wb.remove(wb["Sources"])
src = wb.create_sheet("Sources")
rows = [
    ("Cell", "What it states", "Source", "URL", "Published", "Accessed", "Status"),
    ("assumptions!E2", "Vendor-stated 1-3% recoverable", "Freight audit vendor marketing",
     "see business/data-and-sources.md, invoice error rates", "various", ACCESSED,
     "VENDOR CLAIM, not verified; used only as an assumption"),
    ("assumptions!E3", "15-30% contingency typical", "Audit-vendor pricing pages",
     "see business/data-and-sources.md, competitor pricing", "various", ACCESSED,
     "ASSUMPTION; freight audit firms rarely publish prices"),
    ("assumptions!E8", "Private-SaaS CAC payback median ~18-20 months", "KeyBanc / Benchmarkit",
     "https://www.benchmarkit.ai/2025benchmarks", "2025", ACCESSED, "Not re-verified; secondary"),
    ("valuation!B7", "Public cloud median ~4.1x EV/NTM (Jan 2026)", "Clouded Judgement",
     "https://cloudedjudgement.substack.com/", "2026-01", ACCESSED,
     "Dated; not re-verified for this workbook"),
    ("Source review!B10", "ATRI detention $15.1B, 2023", "ATRI",
     "https://truckingresearch.org/2024/09/new-research-documents-substantial-financial-and-safety-impacts-from-truck-driver-detention/",
     "2024-09-10", ACCESSED, "VERIFIED primary; cost of waiting, not a recovery pool"),
    ("Source review!B8", "FMC D&D ~$15.4B, 9 carriers, Apr 2020-Mar 2025",
     "Federal Maritime Commission", "https://www.fmc.gov/detention-and-demurrage/", "2025",
     ACCESSED, "VERIFIED primary"),
    ("Source review!B9", "46 CFR Part 541; 541.4 reserved", "eCFR / Federal Register 2025-23920",
     "https://www.ecfr.gov/current/title-46/chapter-IV/subchapter-B/part-541",
     "current through 2026-10", ACCESSED, "VERIFIED primary"),
]
for r in rows:
    src.append(r)
for c in src[1]:
    c.font = Font(bold=True)
for col, w in zip("ABCDEFG", (18, 42, 30, 70, 14, 12, 40)):
    src.column_dimensions[col].width = w
wb.save(WB)
for item in log:
    print(*item, sep=" | ")
